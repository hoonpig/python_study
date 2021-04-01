from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference

wb = load_workbook("sample.xlsx")
ws = wb.active

#B2:C11 까지 데이터
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart()          # 차트 종류 설정
# bar_chart.add_data(bar_value)   # 차트 데이터 추가

# ws.add_chart(bar_chart, "E1") # 차트 위치 정의

# 제목을 포함한 B1:C11
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()          # 차트 종류 설정
line_chart.add_data(line_value, titles_from_data=True)   # 차트 데이터 추가 (제목에서 계열 추가)
line_chart.title = "성적표"
line_chart.style = 10 # 미리 정의된 스타일 적용
line_chart.y_axis.title="점수"
line_chart.x_axis.title="번호"

ws.add_chart(line_chart, "E1") # 차트 위치 정의

wb.save("sample_Line_chart.xlsx")


